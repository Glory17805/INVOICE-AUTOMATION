"""Fixes for the gap register, pinned.

Each test names the finding it closes. The point of pinning them is that most
of these gaps are invisible when the system is working: an unbacked-up data
directory looks fine until it isn't, and a reconciliation block that covers
rows 8-50 looks fine for the first forty-three invoices of every month.
"""

from __future__ import annotations

import sqlite3
import zipfile
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app import backup, pipeline, uploads, workbook
from app.config import IRA_INNOVATIONS
from app.gst import rules
from app.models import DocumentType, ExtractedInvoice

PNG = bytes.fromhex("89504e470d0a1a0a") + b"\x00\x00\x00\rIHDR" + b"\x00" * 20
PDF = b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n1 0 obj\n<<>>\nendobj\n"
JPEG = b"\xff\xd8\xff\xe0\x00\x10JFIF" + b"\x00" * 16


# --------------------------------------------------------------------------- #
# S6 - uploads were trusted on their extension
# --------------------------------------------------------------------------- #

@pytest.mark.parametrize("data,expected", [
    (PDF, "pdf"), (PNG, "png"), (JPEG, "jpeg"),
    (b"GIF89a" + b"\x00" * 10, "gif"),
    (b"RIFF\x00\x00\x00\x00WEBPVP8 ", "webp"),
    (b"invoice,amount\n1,200\n", "text"),
    (b"MZ\x00\x00\x90\x00binary", "unknown"),
    (b"", "unknown"),
])
def test_a_file_is_identified_by_its_bytes(data, expected):
    assert uploads.sniff(data).kind == expected


def test_an_extension_that_disagrees_with_the_content_is_refused():
    with pytest.raises(uploads.RejectedUpload, match=r"named \.pdf but the contents are PNG"):
        uploads.verify(PNG, "invoice.pdf", ".pdf")


def test_a_matching_file_passes():
    assert uploads.verify(PDF, "invoice.pdf", ".pdf").kind == "pdf"


def test_an_executable_wearing_a_pdf_name_does_not_reach_the_parser(tmp_path, monkeypatch):
    monkeypatch.setattr(pipeline, "INCOMING_DIR", tmp_path)
    with pytest.raises(ValueError, match="not a readable PDF"):
        pipeline.capture(b"MZ\x00\x00this is a program", "totally-an-invoice.pdf")


# --------------------------------------------------------------------------- #
# D1 - nothing backed up the data directory
# --------------------------------------------------------------------------- #

@pytest.fixture
def scratch_backups(tmp_path, monkeypatch):
    monkeypatch.setattr(backup, "backup_dir", lambda: tmp_path / "backups")
    monkeypatch.setattr(backup, "WORKBOOK_DIR", tmp_path / "workbook")
    monkeypatch.setattr(backup, "ARCHIVE_DIR", tmp_path / "archive")
    (tmp_path / "workbook").mkdir()
    (tmp_path / "archive" / "Jul-26").mkdir(parents=True)
    (tmp_path / "workbook" / "gst-Jul-26.xlsx").write_bytes(b"workbook bytes")
    (tmp_path / "archive" / "Jul-26" / "row0008.pdf").write_bytes(PDF)
    return tmp_path


def test_a_backup_contains_the_irreplaceable_things(scratch_backups):
    written = backup.create(note="test")
    with zipfile.ZipFile(written) as archive:
        names = set(archive.namelist())
    assert "manifest.json" in names
    assert "workbooks/gst-Jul-26.xlsx" in names
    assert "archive/Jul-26/row0008.pdf" in names


def test_a_backup_is_verified_before_it_is_reported_as_written(scratch_backups):
    written = backup.create()
    manifest = backup.verify(written)
    assert manifest["counts"]["workbooks"] == 1

    # A corrupt archive must fail loudly. One that reports success is worse
    # than none, because it gets trusted.
    written.write_bytes(b"not a zip at all")
    with pytest.raises((RuntimeError, zipfile.BadZipFile)):
        backup.verify(written)


def test_the_database_is_snapshotted_through_sqlite_not_copied(scratch_backups, monkeypatch, tmp_path):
    live = tmp_path / "store.db"
    connection = sqlite3.connect(str(live))
    connection.execute("CREATE TABLE documents (id TEXT)")
    connection.execute("INSERT INTO documents VALUES ('doc-1')")
    connection.commit()
    # Left open on purpose: a byte-for-byte copy of a database being written to
    # is how you get an archive that restores corrupt.
    monkeypatch.setattr(backup.db, "path", lambda: live)

    written = backup.create()
    unpacked = backup.restore(written, tmp_path / "restored")
    restored = sqlite3.connect(str(unpacked / "database" / "store.db"))
    assert restored.execute("SELECT id FROM documents").fetchone()[0] == "doc-1"
    restored.close()
    connection.close()


def test_old_archives_are_pruned_but_the_newest_are_kept(scratch_backups, monkeypatch):
    monkeypatch.setattr(backup, "keep_count", lambda: 3)
    for n in range(5):
        (backup.backup_dir()).mkdir(parents=True, exist_ok=True)
        (backup.backup_dir() / f"gst-backup-2026010{n}T000000Z.zip").write_bytes(b"x")
    backup.prune()
    kept = sorted(p.name for p in backup.backup_dir().glob("gst-backup-*.zip"))
    assert len(kept) == 3
    assert kept[-1].endswith("20260104T000000Z.zip")


def test_restoring_never_overwrites_the_live_directory(scratch_backups, tmp_path):
    written = backup.create()
    where = backup.restore(written)
    assert where.exists()
    # The live workbook directory is untouched by a restore.
    assert (tmp_path / "workbook" / "gst-Jul-26.xlsx").read_bytes() == b"workbook bytes"


# --------------------------------------------------------------------------- #
# G1 - the GSTR-1 reconciliation covered only rows 8-50
# --------------------------------------------------------------------------- #

@pytest.fixture
def scratch_books(tmp_path, monkeypatch):
    monkeypatch.setattr(workbook, "WORKBOOK_DIR", tmp_path)
    return tmp_path


def _sale(number, gstin=None, place="Andhra Pradesh", amount=1000):
    return ExtractedInvoice(
        document_type=None, invoice_number=number, invoice_date="15-May-2026",
        supplier_name="Ira Innovations", supplier_gstin=IRA_INNOVATIONS.gstin,
        supplier_address=None, recipient_name="Buyer", recipient_gstin=gstin,
        recipient_address=None, place_of_supply=place, reverse_charge=False,
        is_credit_note=False, hsn_sac=None, line_items=[], taxable_value=amount,
        gst_rate_percent=18, cgst_amount=None, sgst_amount=None, igst_amount=None,
        cess_amount=None, total_amount=None, quantity=None, notes=None)


def _post(doc, period="May-26"):
    treatment = rules.apply_gst(doc, IRA_INNOVATIONS)
    return workbook.post_row(treatment, {
        "period": period, "invoice_number": doc.invoice_number,
        "invoice_date_obj": workbook.parse_date(doc.invoice_date),
        "hsn_sac": None, "quantity": None, "unit_rate": None}, period)


def test_the_b2b_b2c_split_is_derived_not_a_hand_written_row_list(scratch_books):
    workbook.ensure_working_copy("May-26")
    _post(_sale("S-1", gstin="29AAFCB7707D1ZQ", place="Karnataka"))

    ws = load_workbook(workbook.workbook_path("May-26"))["GSTR-1"]
    b2b, b2c = workbook._reconciliation_rows(ws)
    # The old sheet had "=J8+J34" here, which is a list of two rows.
    assert ws[f"J{b2b}"].value.startswith("=SUMIF(")
    assert "$D$8:$D$" in ws[f"J{b2b}"].value
    assert ws[f"J{b2c}"].value.startswith("=SUM(")


def test_the_split_covers_every_row_past_the_old_fifty_row_limit(scratch_books):
    """60 invoices: the old formula stopped at row 50 and silently under-reported."""
    workbook.ensure_working_copy("May-26")
    for n in range(60):
        b2b = n % 2 == 0
        _post(_sale(f"S-{n:03d}",
                    gstin="29AAFCB7707D1ZQ" if b2b else None,
                    place="Karnataka" if b2b else "Andhra Pradesh"))

    ws = load_workbook(workbook.workbook_path("May-26"))["GSTR-1"]
    spec = workbook.SPECS[DocumentType.SALES]
    totals = workbook.find_totals_row(ws, spec)
    last = totals - 1
    b2b_row, _ = workbook._reconciliation_rows(ws)

    # The formula's range must reach the final data row, not stop at 50.
    assert f"$D$8:$D${last}" in ws[f"J{b2b_row}"].value
    assert last >= 60

    # Evaluate it the way Excel would: B2B is every row carrying a GSTIN.
    b2b_total = sum(ws[f"J{r}"].value or 0 for r in range(8, last + 1)
                    if ws[f"D{r}"].value)
    register = sum(ws[f"J{r}"].value or 0 for r in range(8, last + 1))
    assert b2b_total == pytest.approx(30_000)
    assert register - b2b_total == pytest.approx(30_000)


def test_the_block_is_found_by_label_so_growing_the_register_cannot_orphan_it(scratch_books):
    """Inserting rows pushes the block down; a hard-coded row 75 would break."""
    workbook.ensure_working_copy("May-26")
    ws = load_workbook(workbook.workbook_path("May-26"))["GSTR-1"]
    before = workbook._reconciliation_rows(ws)

    spec = workbook.SPECS[DocumentType.SALES]
    template_rows = workbook.find_totals_row(ws, spec) - spec.first_data_row
    for n in range(template_rows + 2):        # two rows past the template
        _post(_sale(f"G-{n:03d}"))

    ws = load_workbook(workbook.workbook_path("May-26"))["GSTR-1"]
    after = workbook._reconciliation_rows(ws)
    assert after is not None and after[0] > before[0], "block should have moved down"
    assert ws[f"J{after[0]}"].value.startswith("=SUMIF(")


# --------------------------------------------------------------------------- #
# G2 - opening credit was never carried forward
# --------------------------------------------------------------------------- #

def test_a_new_period_inherits_the_previous_periods_closing_credit(scratch_books):
    workbook.ensure_working_copy("May-26")
    may = workbook.tax_payable_summary("May-26")

    workbook.ensure_working_copy("Jun-26")
    june = workbook.tax_payable_summary("Jun-26")

    for key in ("igst", "cgst", "sgst"):
        closing = Decimal(may.itc_available[key]) - Decimal(may.output_tax[key])
        assert Decimal(june.itc_carry_forward[key]) == max(closing, Decimal("0"))
    assert not workbook.opening_credit_is_unset("Jun-26")


def test_a_period_with_no_predecessor_still_starts_at_zero_and_says_so(scratch_books):
    """Inventing an opening credit overstates it, which understates tax due."""
    workbook.ensure_working_copy("Dec-30")
    summary = workbook.tax_payable_summary("Dec-30")
    assert summary.itc_carry_forward == {"igst": "0.00", "cgst": "0.00", "sgst": "0.00"}
    assert workbook.opening_credit_is_unset("Dec-30")


# --------------------------------------------------------------------------- #
# D2 - money crossed the wire as floating point
# --------------------------------------------------------------------------- #

def test_tax_figures_leave_as_exact_decimal_strings(scratch_books):
    workbook.ensure_working_copy("May-26")
    summary = workbook.tax_payable_summary("May-26")
    for group in (summary.itc_available, summary.output_tax, summary.net_payable):
        for value in group.values():
            assert isinstance(value, str)
            # Two decimal places, always - never 616509.0000000001.
            assert Decimal(value) == Decimal(value).quantize(Decimal("0.01"))


def test_a_figure_a_float_cannot_hold_survives_the_round_trip(scratch_books):
    workbook.ensure_working_copy("May-26")
    _post(_sale("PREC-1", amount=4881.10))
    summary = workbook.tax_payable_summary("May-26")
    # 4881.10 at 18%, halved, is 439.299 -> 439.30 exactly.
    assert summary.output_tax["cgst"] == "439.30"
    assert summary.output_tax["sgst"] == "439.30"


# --------------------------------------------------------------------------- #
# A5 - unhandled errors, and S4 - security headers
#
# These live behind a route that only runs when something has already gone
# wrong, which is exactly why nothing exercised them: the first version of the
# handler referred to JSONResponse without importing it, so the one thing
# standing between a traceback and a browser would itself have raised
# NameError. Importing the module could not catch that - the name only resolves
# when the handler is called - so it has to be called.
# --------------------------------------------------------------------------- #

@pytest.fixture
def client_with_a_broken_route():
    from fastapi.testclient import TestClient

    from app.main import app

    @app.get("/api/_test_explode")
    def explode():
        raise RuntimeError("a very specific internal detail")

    # raise_server_exceptions=False makes the client behave like a browser:
    # it returns the 500 the handler produced instead of re-raising.
    try:
        yield TestClient(app, raise_server_exceptions=False)
    finally:
        app.router.routes[:] = [
            r for r in app.router.routes
            if getattr(r, "path", None) != "/api/_test_explode"
        ]


def test_an_unhandled_error_returns_a_reference_and_leaks_nothing(client_with_a_broken_route):
    response = client_with_a_broken_route.get("/api/_test_explode")
    assert response.status_code == 500

    body = response.json()
    assert "reference" in body and len(body["reference"]) == 8
    assert body["reference"] in body["detail"]
    # It must say the request did not complete - an accountant mid-filing needs
    # to know whether their invoice posted.
    assert "did not complete" in body["detail"]

    # And it must not hand the internals of a tax system to the browser.
    raw = response.text
    assert "a very specific internal detail" not in raw
    assert "Traceback" not in raw
    assert "RuntimeError" not in raw


def test_every_response_carries_the_security_headers(client_with_a_broken_route):
    """Including the error path, which is the one most likely to be forgotten."""
    for path in ("/api/health", "/api/_test_explode"):
        headers = client_with_a_broken_route.get(path).headers
        assert headers["X-Content-Type-Options"] == "nosniff"
        assert headers["X-Frame-Options"] == "DENY"
        assert headers["Referrer-Policy"] == "no-referrer"
        assert "frame-ancestors 'none'" in headers["Content-Security-Policy"]
        assert headers["Cache-Control"] == "no-store"


def test_hsts_is_not_sent_over_plain_http(client_with_a_broken_route):
    """Sent over http it is ignored anyway, and setting it while developing on
    127.0.0.1 pins a localhost HSTS entry that breaks every other local project
    until the user clears it by hand."""
    assert "Strict-Transport-Security" not in client_with_a_broken_route.get("/api/health").headers


def test_hsts_is_sent_when_a_proxy_reports_tls(client_with_a_broken_route):
    headers = client_with_a_broken_route.get(
        "/api/health", headers={"x-forwarded-proto": "https"}
    ).headers
    assert "max-age=31536000" in headers["Strict-Transport-Security"]


# --------------------------------------------------------------------------- #
# The dashboard feed shows invoice events; the audit trail keeps everything
# --------------------------------------------------------------------------- #

def test_the_dashboard_feed_excludes_sign_in_noise():
    """Six failed sign-ins against four uploads crowded out the thing the feed
    exists to show. Filtering is a view concern - nothing stops being recorded."""
    from app import accounts

    for action in ("login", "login_failed", "logout", "signup",
                   "settings_updated", "profile_updated"):
        assert action not in accounts.DOCUMENT_ACTIONS, action
    for action in ("uploaded", "posted", "document_deleted", "posted_with_override"):
        assert action in accounts.DOCUMENT_ACTIONS, action


def test_the_audit_trail_still_records_and_returns_everything():
    """The compliance property: every action against the workbook is attributable,
    and sign-ins remain visible to an administrator."""
    from app import accounts

    accounts.record(None, "login_failed", {"email": "someone@example.test"})
    accounts.record(None, "uploaded", {"count": 1})

    unfiltered = {row["action"] for row in accounts.activity(200)}
    assert "login_failed" in unfiltered, "the audit trail must not lose sign-in events"
    assert "uploaded" in unfiltered

    filtered = {row["action"] for row in accounts.activity(200, accounts.DOCUMENT_ACTIONS)}
    assert "login_failed" not in filtered
    assert "uploaded" in filtered
