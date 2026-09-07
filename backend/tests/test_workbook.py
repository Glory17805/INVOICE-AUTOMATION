"""Workbook write tests: rows land on the right sheet, in the sheet's own idiom.

These run against a throwaway copy of the real workbook so the assertions are
about the client's actual sheet layout, not a fixture that agrees with us.
"""

from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app import workbook
from app.config import IRA_INNOVATIONS
from app.gst import rules
from app.models import DocumentType

from .test_gst import invoice

# Every fixture invoice below is dated in May 2026, so they file against the
# period the master workbook itself covers.
PERIOD = "May-26"


@pytest.fixture(autouse=True)
def fresh_workbook(tmp_path, monkeypatch):
    """Point the module at a scratch workbook directory for every test."""
    monkeypatch.setattr(workbook, "WORKBOOK_DIR", tmp_path)
    workbook.ensure_working_copy(PERIOD)
    yield workbook.workbook_path(PERIOD)


def post(doc, period=PERIOD):
    treatment = rules.apply_gst(doc, IRA_INNOVATIONS)
    meta = {
        "period": period,
        "invoice_number": doc.invoice_number,
        "invoice_date_obj": workbook.parse_date(doc.invoice_date),
        "hsn_sac": doc.hsn_sac,
        "quantity": doc.quantity,
        "unit_rate": None,
    }
    return workbook.post_row(treatment, meta, period)


def test_source_workbook_is_never_modified(fresh_workbook):
    from app.config import source_workbook
    before = source_workbook().stat().st_mtime_ns
    post(invoice(
        invoice_number="IRA1082026-27", invoice_date="31-May-2026",
        supplier_gstin=IRA_INNOVATIONS.gstin, supplier_name="Ira Innovations",
        recipient_name="Prasuna Reddy", place_of_supply="Andhra Pradesh ** ( 28 )",
        taxable_value=4881.10, gst_rate_percent=18,
    ))
    assert source_workbook().stat().st_mtime_ns == before


def test_a_sale_lands_on_gstr1_using_the_sheets_own_formulas(fresh_workbook):
    sheet, row = post(invoice(
        invoice_number="IRA1082026-27", invoice_date="31-May-2026",
        supplier_gstin=IRA_INNOVATIONS.gstin, supplier_name="Ira Innovations",
        recipient_name="Prasuna Reddy", place_of_supply="Andhra Pradesh ** ( 28 )",
        hsn_sac="48182000", quantity=209, taxable_value=4881.10, gst_rate_percent=18,
    ))
    assert sheet == "GSTR-1"
    assert row == 8  # the first pre-formatted template row

    ws = load_workbook(fresh_workbook)["GSTR-1"]
    assert ws["C8"].value == "IRA1082026-27"
    assert ws["E8"].value == "Prasuna Reddy"
    assert ws["J8"].value == pytest.approx(4881.10)
    # Intra-state: the CGST/SGST formula pattern the sheet already uses.
    assert ws["K8"].value == 0
    assert ws["L8"].value == "=J8*I8/2"
    assert ws["M8"].value == "=L8"
    assert ws["N8"].value == "=J8+K8+L8+M8"


def test_an_inter_state_sale_uses_the_igst_formula_pattern(fresh_workbook):
    _, row = post(invoice(
        invoice_number="IRA9990001", invoice_date="12-May-2026",
        supplier_gstin=IRA_INNOVATIONS.gstin, supplier_name="Ira Innovations",
        recipient_name="Karnataka Buyer", recipient_gstin="29AAFCB7707D1ZQ",
        place_of_supply="Karnataka", taxable_value=10000, gst_rate_percent=18,
    ))
    ws = load_workbook(fresh_workbook)["GSTR-1"]
    assert ws[f"K{row}"].value == f"=I{row}*J{row}"
    assert ws[f"L{row}"].value == 0


def test_a_purchase_appends_below_the_existing_rows_and_grows_the_sum(fresh_workbook):
    sheet, row = post(invoice(
        invoice_number="TEST-PUR-1", invoice_date="15/05/2026",
        supplier_name="Waymiro Private Limited", supplier_gstin="21AADCW9393G1Z3",
        recipient_gstin=IRA_INNOVATIONS.gstin, taxable_value=6000, gst_rate_percent=18,
    ))
    assert sheet == "GSTR-2B"
    assert row == 9  # appended after the workbook's two existing rows

    ws = load_workbook(fresh_workbook)["GSTR-2B"]
    assert ws["C9"].value == "TEST-PUR-1"
    assert ws["E9"].value == "WAYMIRO PRIVATE LIMITED"
    assert ws["H9"].value == pytest.approx(1080.0)   # IGST, inter-state
    assert ws["I9"].value == 0
    # The totals row moved down and its SUM now covers the new row.
    assert ws["G10"].value == "=SUM(G7:G9)"


def test_growing_a_register_repoints_the_tax_payable_formulas(fresh_workbook):
    post(invoice(
        invoice_number="TEST-PUR-2", invoice_date="15-May-2026", supplier_gstin="21AADCW9393G1Z3",
        recipient_gstin=IRA_INNOVATIONS.gstin, taxable_value=1000, gst_rate_percent=18,
    ))
    tp = load_workbook(fresh_workbook)["Tax Payable"]
    # Was H9 before the insert; must now follow the totals row to H10.
    assert tp["E7"].value == "=ROUND('GSTR-2B'!H10,0)"
    assert tp["F7"].value == "=ROUND('GSTR-2B'!I10,0)"


def test_reverse_charge_lands_on_the_rcm_sheet_with_no_cgst_or_sgst(fresh_workbook):
    sheet, row = post(invoice(
        invoice_number="CRN0000001", invoice_date="06/05/2026",
        supplier_name="Porter", supplier_gstin="36AAGCR8772D1Z3",
        recipient_gstin=IRA_INNOVATIONS.gstin, reverse_charge=True,
        taxable_value=682, gst_rate_percent=5,
    ))
    assert sheet == "RCM"
    ws = load_workbook(fresh_workbook)["RCM"]
    assert ws[f"F{row}"].value == "Yes"
    assert ws[f"I{row}"].value == pytest.approx(34.10)
    assert ws[f"J{row}"].value == 0
    assert ws[f"K{row}"].value == 0
    # The RCM total is the bill value; the tax is a separate cash liability.
    assert ws[f"M{row}"].value == pytest.approx(682.0)


def test_a_credit_note_lands_on_the_credit_note_sheet(fresh_workbook):
    sheet, row = post(invoice(
        invoice_number="CN-TEST-1", invoice_date="13/05/2026",
        supplier_name="Swiggy", supplier_gstin="29AAFCB7707D1ZQ",
        recipient_gstin=IRA_INNOVATIONS.gstin, is_credit_note=True,
        place_of_supply="Andhra Pradesh", taxable_value=30000, gst_rate_percent=18,
    ))
    assert sheet == "Credit Note"
    ws = load_workbook(fresh_workbook)["Credit Note"]
    assert ws[f"E{row}"].value == "Credit Note"
    assert ws[f"L{row}"].value == pytest.approx(5400.0)
    assert ws[f"H{row}"].value == "Andhra Pradesh"


def test_posted_rows_come_back_out_of_the_register(fresh_workbook):
    post(invoice(
        invoice_number="IRA1082026-27", invoice_date="31-May-2026",
        supplier_gstin=IRA_INNOVATIONS.gstin, supplier_name="Ira Innovations",
        recipient_name="Prasuna Reddy", place_of_supply="Andhra Pradesh",
        taxable_value=4881.10, gst_rate_percent=18,
    ))
    rows = workbook.read_register(DocumentType.SALES, PERIOD)
    assert len(rows) == 1
    values = rows[0]["values"]
    assert values["Invoice no"] == "IRA1082026-27"
    # Formulas are evaluated the way the sheet would evaluate them.
    assert values["CGST"] == "439.30"
    assert values["SGST"] == "439.30"
    assert values["Invoice Amount"] == "5,759.70"


def test_duplicate_detection_reads_back_from_the_register(fresh_workbook):
    post(invoice(
        invoice_number="IRA1082026-27", invoice_date="31-May-2026", supplier_gstin=IRA_INNOVATIONS.gstin,
        supplier_name="Ira Innovations", recipient_name="Prasuna Reddy",
        place_of_supply="Andhra Pradesh", taxable_value=4881.10, gst_rate_percent=18,
    ))
    assert ("IRA1082026-27", None) in workbook.posted_keys(DocumentType.SALES, PERIOD)


def test_unposting_clears_the_row_and_restores_the_template(fresh_workbook):
    sheet, row = post(invoice(
        invoice_number="IRA1082026-27", invoice_date="31-May-2026", supplier_gstin=IRA_INNOVATIONS.gstin,
        supplier_name="Ira Innovations", recipient_name="Prasuna Reddy",
        place_of_supply="Andhra Pradesh", taxable_value=4881.10, gst_rate_percent=18,
    ))
    workbook.unpost_row(sheet, row, PERIOD)
    assert workbook.read_register(DocumentType.SALES, PERIOD) == []
    ws = load_workbook(fresh_workbook)["GSTR-1"]
    assert ws[f"L{row}"].value == f"=J{row}*I{row}/2"


def test_tax_payable_reflects_the_workbooks_own_opening_position(fresh_workbook):
    summary = workbook.tax_payable_summary(PERIOD)
    # ITC carried forward, straight off the Tax Payable sheet.
    assert Decimal(summary.itc_carry_forward["igst"]) == Decimal("620777")
    assert Decimal(summary.itc_carry_forward["cgst"]) == Decimal("33262")
    # Opening ITC available matches the workbook's own May-26 figures.
    assert Decimal(summary.itc_available["igst"]) == Decimal("616509")
    assert Decimal(summary.itc_available["cgst"]) == Decimal("33354")
    assert Decimal(summary.itc_available["sgst"]) == Decimal("33354")
    # No sales entered yet, so no output tax.
    assert Decimal(summary.output_tax["cgst"]) == Decimal("0")


def test_posting_a_sale_moves_the_output_tax_position(fresh_workbook):
    post(invoice(
        invoice_number="IRA1082026-27", invoice_date="31-May-2026", supplier_gstin=IRA_INNOVATIONS.gstin,
        supplier_name="Ira Innovations", recipient_name="Prasuna Reddy",
        place_of_supply="Andhra Pradesh", taxable_value=4881.10, gst_rate_percent=18,
    ))
    summary = workbook.tax_payable_summary(PERIOD)
    assert Decimal(summary.output_tax["cgst"]) == Decimal("439.30")
    assert Decimal(summary.output_tax["sgst"]) == Decimal("439.30")
    # Credit on hand far exceeds it, so nothing is payable in cash.
    assert Decimal(summary.net_payable["cgst"]) == Decimal("0")


def test_gstr1_grows_past_its_last_template_row(fresh_workbook):
    """The sheet ships with 64 rows; posting a 65th must extend it, not fail."""
    spec = workbook.SPECS[DocumentType.SALES]
    wb = load_workbook(fresh_workbook)
    template_rows = workbook.find_totals_row(wb["GSTR-1"], spec) - spec.first_data_row
    wb.close()

    for n in range(template_rows + 1):
        post(invoice(
            invoice_number=f"BULK-{n:03d}", invoice_date="20-May-2026", supplier_gstin=IRA_INNOVATIONS.gstin,
            supplier_name="Ira Innovations", recipient_name="Prasuna Reddy",
            place_of_supply="Andhra Pradesh", taxable_value=100, gst_rate_percent=18,
        ))

    rows = workbook.read_register(DocumentType.SALES, PERIOD)
    assert len(rows) == template_rows + 1

    ws = load_workbook(fresh_workbook)["GSTR-1"]
    totals = workbook.find_totals_row(ws, spec)
    assert totals == spec.first_data_row + template_rows + 1
    assert ws[f"J{totals}"].value == f"=SUM(J8:J{totals - 1})"


def test_sales_totals_reconcile_with_the_sheet_sum(fresh_workbook):
    for n, amount in enumerate([1000, 2500.55, 700.25]):
        post(invoice(
            invoice_number=f"SUM-{n}", invoice_date="20-May-2026", supplier_gstin=IRA_INNOVATIONS.gstin,
            supplier_name="Ira Innovations", recipient_name="Prasuna Reddy",
            place_of_supply="Andhra Pradesh", taxable_value=amount, gst_rate_percent=18,
        ))
    summary = workbook.tax_payable_summary(PERIOD)
    expected = Decimal("0")
    for amount in [Decimal("1000"), Decimal("2500.55"), Decimal("700.25")]:
        expected += (amount * Decimal("0.18") / 2).quantize(Decimal("0.01"))
    assert Decimal(summary.output_tax["cgst"]) == expected


# --------------------------------------------------------------------------- #
# Display: a posted row has to be readable in Excel, not just correct
# --------------------------------------------------------------------------- #

def test_posted_rows_keep_the_sheets_own_number_formats(fresh_workbook):
    """A wider format does not widen the column - Excel renders '####' instead.

    The rate column is 5.1 characters wide because '0%' fits it, and the filing
    period column is 7.3 wide because it shows 'Apr-26'. Imposing '0.00%' or a
    full date format makes a correct number invisible, so posted rows must adopt
    the format already in the column.
    """
    post(invoice(
        invoice_number="IRA1082026-27", invoice_date="31-May-2026",
        supplier_gstin=IRA_INNOVATIONS.gstin, supplier_name="Ira Innovations",
        recipient_name="Prasuna Reddy", place_of_supply="Andhra Pradesh ** ( 28 )",
        taxable_value=4881.10, gst_rate_percent=18,
    ))
    post(invoice(
        invoice_number="PUR-1", invoice_date="15-May-2026",
        supplier_name="Waymiro Private Limited", supplier_gstin="21AADCW9393G1Z3",
        recipient_gstin=IRA_INNOVATIONS.gstin, taxable_value=6000, gst_rate_percent=18,
    ))
    post(invoice(
        invoice_number="RCM-1", invoice_date="06-May-2026", supplier_name="Porter",
        supplier_gstin="36AAGCR8772D1Z3", recipient_gstin=IRA_INNOVATIONS.gstin,
        reverse_charge=True, taxable_value=682, gst_rate_percent=5,
    ))

    wb = load_workbook(fresh_workbook)
    # GSTR-1: the sheet's own date and percent formats, not wider substitutes.
    assert wb["GSTR-1"]["B8"].number_format == "d-mmm-yy"
    assert wb["GSTR-1"]["I8"].number_format == "0%"
    # An appended row inherits the format its neighbours use. Assigning a
    # datetime makes openpyxl stamp its own wide format, which must not survive.
    for sheet, cell in (("GSTR-2B", "A9"), ("RCM", "A11")):
        assert wb[sheet][cell].number_format == "mmm-yy", f"{sheet}!{cell}"


def test_a_grown_gstr1_row_inherits_the_template_formats(fresh_workbook):
    """Rows past the last template row must format like the rows above them."""
    spec = workbook.SPECS[DocumentType.SALES]
    wb = load_workbook(fresh_workbook)
    template_rows = workbook.find_totals_row(wb["GSTR-1"], spec) - spec.first_data_row
    wb.close()

    for n in range(template_rows + 1):
        post(invoice(
            invoice_number=f"GROW-{n:03d}", invoice_date="31-May-2026",
            supplier_gstin=IRA_INNOVATIONS.gstin, supplier_name="Ira Innovations",
            recipient_name="Prasuna Reddy", place_of_supply="Andhra Pradesh",
            taxable_value=100, gst_rate_percent=18,
        ))

    grown = spec.first_data_row + template_rows   # the row that did not exist before
    ws = load_workbook(fresh_workbook)["GSTR-1"]
    assert ws[f"B{grown}"].number_format == "d-mmm-yy"
    assert ws[f"I{grown}"].number_format == "0%"


# --------------------------------------------------------------------------- #
# Reading back rows a person typed into Excel by hand
#
# The client has kept these sheets by hand for years and will keep doing it
# alongside the app. A typed row carries plain numbers where a posted row
# carries formulas, and both have to read back as what they actually say.
# --------------------------------------------------------------------------- #

def _type_row_by_hand(path, row: int, **cells) -> None:
    """Write literal values into GSTR-1 the way a person entering data would."""
    wb = load_workbook(path)
    ws = wb["GSTR-1"]
    for column, value in cells.items():
        ws[f"{column}{row}"] = value
    wb.save(path)
    wb.close()


def test_a_hand_typed_inter_state_sale_is_not_read_as_cgst_and_sgst_too(fresh_workbook):
    """The defect this guards: inferring the split from one cell's *type*.

    A posted inter-state row puts a formula in K. A person typing the same sale
    puts a number there. Treating "K is not a formula" as "this row must be
    CGST + SGST" invents a tax that is not on the sheet - here it would add
    900 + 900 to a row whose only tax is 1800 of IGST.
    """
    _type_row_by_hand(
        fresh_workbook, 8,
        C="HAND-INTER-1", E="Karnataka Buyer", D="29AAFCB7707D1ZQ",
        I=0.18, J=10000, K=1800, L=None, M=None, N=11800,
    )

    rows = workbook.read_register(DocumentType.SALES, PERIOD)
    assert len(rows) == 1
    values = rows[0]["values"]
    assert values["IGST"] == "1,800.00"
    assert values["CGST"] == "0.00"
    assert values["SGST"] == "0.00"
    assert values["Invoice Amount"] == "11,800.00"

    # And it must not inflate the position the Tax Payable screen reports.
    summary = workbook.tax_payable_summary(PERIOD)
    assert Decimal(summary.output_tax["igst"]) == Decimal("1800.0")
    assert Decimal(summary.output_tax["cgst"]) == Decimal("0.0")
    assert Decimal(summary.output_tax["sgst"]) == Decimal("0.0")


def test_a_hand_typed_intra_state_sale_reads_back_as_typed(fresh_workbook):
    """The complement: typed CGST/SGST is taken at face value, not recomputed."""
    _type_row_by_hand(
        fresh_workbook, 8,
        C="HAND-INTRA-1", E="Prasuna Reddy",
        I=0.18, J=10000, K=0, L=900, M=900, N=11800,
    )

    values = workbook.read_register(DocumentType.SALES, PERIOD)[0]["values"]
    assert values["IGST"] == "0.00"
    assert values["CGST"] == "900.00"
    assert values["SGST"] == "900.00"
    assert values["Invoice Amount"] == "11,800.00"


def test_the_sheets_own_formula_idioms_still_evaluate(fresh_workbook):
    """Posted rows are unaffected: both formula patterns evaluate as before."""
    # Intra-state, as _write_sales lays it out: K=0, L==J*I/2, M==L.
    post(invoice(
        invoice_number="FORMULA-INTRA", invoice_date="31-May-2026",
        supplier_gstin=IRA_INNOVATIONS.gstin, supplier_name="Ira Innovations",
        recipient_name="Prasuna Reddy", place_of_supply="Andhra Pradesh",
        taxable_value=4881.10, gst_rate_percent=18,
    ))
    # Inter-state: K==I*J, L=0, M==L.
    post(invoice(
        invoice_number="FORMULA-INTER", invoice_date="12-May-2026",
        supplier_gstin=IRA_INNOVATIONS.gstin, supplier_name="Ira Innovations",
        recipient_name="Karnataka Buyer", recipient_gstin="29AAFCB7707D1ZQ",
        place_of_supply="Karnataka", taxable_value=10000, gst_rate_percent=18,
    ))

    by_number = {r["values"]["Invoice no"]: r["values"] for r in
                 workbook.read_register(DocumentType.SALES, PERIOD)}

    intra = by_number["FORMULA-INTRA"]
    assert (intra["CGST"], intra["SGST"], intra["IGST"]) == ("439.30", "439.30", "0.00")

    inter = by_number["FORMULA-INTER"]
    assert (inter["IGST"], inter["CGST"], inter["SGST"]) == ("1,800.00", "0.00", "0.00")


def test_an_unreadable_formula_scores_zero_rather_than_guessing(fresh_workbook):
    """openpyxl reads formulas, not results. A formula we cannot evaluate is
    reported as nothing rather than as an invented number."""
    _type_row_by_hand(
        fresh_workbook, 8,
        C="ODD-1", E="Someone", I=0.18, J=10000, K="=SOMETHING_ELSE(1)", L=0, M=0,
    )
    values = workbook.read_register(DocumentType.SALES, PERIOD)[0]["values"]
    assert values["IGST"] == "0.00"


# --------------------------------------------------------------------------- #
# Downloading every period at once
# --------------------------------------------------------------------------- #

def test_all_periods_download_as_one_file_per_period(fresh_workbook, monkeypatch):
    """The bundle holds each period separately, not merged into one workbook.

    Merging would produce a file matching no return that can be filed: every
    period carries its own totals and its own tax position.
    """
    import io
    import zipfile

    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from app import accounts, main

    workbook.ensure_working_copy("Jun-26")
    workbook.ensure_working_copy("Jul-26")

    recorded = []
    monkeypatch.setattr(accounts, "record",
                        lambda user, action, detail=None: recorded.append((action, detail)))

    app = FastAPI()
    app.get("/api/workbook/download-all")(
        lambda: main.download_all_workbooks(user={"id": "u", "email": "a@b.com"}))

    response = TestClient(app).get("/api/workbook/download-all")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/zip"

    with zipfile.ZipFile(io.BytesIO(response.content)) as bundle:
        names = sorted(bundle.namelist())
        assert names == sorted(f"Ira Innovations GST {p}.xlsx"
                               for p in workbook.available_periods())
        # Every entry is a real workbook, not an empty placeholder.
        for name in names:
            with bundle.open(name) as member:
                assert load_workbook(io.BytesIO(member.read())).sheetnames

    assert recorded and recorded[0][0] == "workbook_downloaded"
    assert recorded[0][1]["period"] == "all"


def test_all_periods_download_refuses_when_there_is_nothing(monkeypatch):
    from fastapi import HTTPException

    from app import main

    monkeypatch.setattr(workbook, "available_periods", list)
    with pytest.raises(HTTPException) as raised:
        main.download_all_workbooks(user={"id": "u", "email": "a@b.com"})
    assert raised.value.status_code == 404
